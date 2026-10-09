"""
build_ground_truth.py
=====================
Chuyển đổi file ground_truth_master.json sang file Excel GT.xlsx chuẩn format của eval_metrics.py:
- Sheet 'ground_truth': doc_id | field_key | value
- Sheet 'table_<key>': doc_id | row | <cols...>
"""
from __future__ import annotations

import json
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

JSON_PATH = Path("/Users/apple/ocr-pipeline/ground_truth_master.json")
OUTPUT_PATH = Path("/Users/apple/ocr-pipeline/GT.xlsx")

def main():
    with open(JSON_PATH, "r", encoding="utf-8") as f:
        master = json.load(f)

    wb = openpyxl.Workbook()
    ws_gt = wb.active
    ws_gt.title = "ground_truth"

    fill_header = PatternFill("solid", fgColor="4F46E5")
    font_header = Font(bold=True, color="FFFFFF")

    headers = ["doc_id", "field_key", "value"]
    ws_gt.append(headers)
    for c in ws_gt[1]:
        c.fill = fill_header
        c.font = font_header
        c.alignment = Alignment(horizontal="center", vertical="center")

    table_data = {}

    for doc in master.get("tai_lieu", []):
        doc_id = doc.get("ma")
        fields = doc.get("truong_du_lieu", {})

        def add_field(key, val):
            if val is not None and not isinstance(val, (dict, list)):
                ws_gt.append([doc_id, key, str(val)])

        # Xử lý theo từng tài liệu
        if doc_id == "B1":
            add_field("so_hop_dong", fields.get("so_hop_dong"))
            add_field("ngay_ky", fields.get("ngay_ky"))
            add_field("dia_diem_ky", fields.get("dia_diem_ky"))
            add_field("so_ban_hop_dong", fields.get("so_ban_hop_dong"))
            
            # bên nhận
            bn = fields.get("ben_nhan_the_chap", {})
            add_field("ben_nhan_ten", bn.get("ten"))
            add_field("ben_nhan_dia_chi", bn.get("dia_chi"))
            add_field("ben_nhan_ma_so_thue", bn.get("ma_so_thue"))
            add_field("ben_nhan_nguoi_dai_dien", bn.get("nguoi_dai_dien"))
            add_field("ben_nhan_chuc_vu", bn.get("chuc_vu"))

            # tài sản
            ts = fields.get("tai_san_the_chap", {})
            add_field("dia_chi_tai_san", ts.get("dia_chi_tai_san"))
            add_field("thua_dat_so", ts.get("thua_dat_so"))
            add_field("to_ban_do_so", ts.get("to_ban_do_so"))
            add_field("dien_tich_dat_m2", ts.get("dien_tich_dat_m2"))
            add_field("muc_dich_su_dung", ts.get("muc_dich_su_dung"))
            add_field("thoi_han_su_dung", ts.get("thoi_han_su_dung"))
            add_field("so_phat_hanh_gcn", ts.get("so_phat_hanh_gcn"))
            add_field("so_vao_so_gcn", ts.get("so_vao_so_gcn"))
            add_field("ngay_cap_gcn", ts.get("ngay_cap_gcn"))
            add_field("co_quan_cap_gcn", ts.get("co_quan_cap_gcn"))

            # nhà
            nha = fields.get("tai_san_gan_lien_voi_dat", {})
            add_field("loai_nha", nha.get("loai"))
            add_field("dien_tich_xay_dung", nha.get("dien_tich_xay_dung"))
            add_field("tong_dien_tich_san", nha.get("tong_dien_tich_san"))

            # định giá
            dg = fields.get("gia_tri_dinh_gia", {})
            add_field("gia_tri_dinh_gia_bang_so", dg.get("tong_bang_so"))
            add_field("gia_tri_dinh_gia_bang_chu", dg.get("tong_bang_chu"))
            add_field("gia_tri_quyen_su_dung_dat", dg.get("gia_tri_quyen_su_dung_dat"))
            add_field("gia_tri_tai_san_gan_lien", dg.get("gia_tri_tai_san_gan_lien"))
            add_field("ngay_dinh_gia", dg.get("ngay_dinh_gia"))
            add_field("don_vi_dinh_gia", dg.get("don_vi_dinh_gia"))

            # nghĩa vụ
            nv = fields.get("nghia_vu_duoc_bao_dam", {})
            add_field("so_hop_dong_tin_dung", nv.get("so_hop_dong_tin_dung"))
            add_field("ngay_hop_dong_tin_dung", nv.get("ngay_hop_dong_tin_dung"))
            add_field("so_tien_vay_bang_so", nv.get("so_tien_vay_bang_so"))
            add_field("so_tien_vay_bang_chu", nv.get("so_tien_vay_bang_chu"))
            add_field("muc_dich_vay", nv.get("muc_dich_vay"))
            add_field("thoi_han_vay", nv.get("thoi_han_vay"))

            # công chứng
            cc = fields.get("cong_chung", {})
            add_field("to_chuc_cong_chung", cc.get("to_chuc_cong_chung"))
            add_field("cong_chung_vien", cc.get("cong_chung_vien"))
            add_field("so_cong_chung", cc.get("so_cong_chung"))
            add_field("quyen_so", cc.get("quyen_so"))
            add_field("ngay_cong_chung", cc.get("ngay_cong_chung"))
            add_field("so_dieu_khoan", fields.get("so_dieu_khoan"))

            # lịch trả nợ
            lt = fields.get("lich_tra_no", {})
            add_field("lich_tra_no_ky_dau_ngay", lt.get("ky_dau_ngay"))
            add_field("lich_tra_no_ky_cuoi_ngay", lt.get("ky_cuoi_ngay"))
            add_field("lich_tra_no_tong_goc", lt.get("tong_goc"))

            # bảng ben_the_chap
            table_data.setdefault("table_ben_the_chap", []).extend([
                (doc_id, idx + 1, item.get("ho_ten"), item.get("ngay_sinh"), item.get("so_cccd"), item.get("ngay_cap_cccd"), item.get("noi_thuong_tru"))
                for idx, item in enumerate(fields.get("ben_the_chap", []))
            ])

        elif doc_id == "B2":
            add_field("co_quan_dang_ky", fields.get("co_quan_nhan"))
            add_field("so_van_ban", fields.get("phan_ghi_cua_co_quan_dang_ky", {}).get("so_vao_so_tiep_nhan"))
            add_field("loai_ben_bao_dam", "Cá nhân")
            
            btc_list = fields.get("ben_the_chap", [])
            if btc_list:
                add_field("ten_ben_bao_dam", btc_list[0].get("ho_ten"))
                add_field("msdn_msdp", btc_list[0].get("so_cccd"))
                add_field("dia_chi_tru_so", btc_list[0].get("noi_thuong_tru"))
            
            bn = fields.get("ben_nhan_the_chap", {})
            add_field("ten_ben_nhan_bao_dam", bn.get("ten"))
            add_field("dia_chi_ben_nhan", bn.get("dia_chi"))
            add_field("ngay_dang_ky", fields.get("hop_dong_bao_dam", {}).get("ngay_ky"))

            ts = fields.get("tai_san_bao_dam", {})
            add_field("loai_tai_san_bao_dam", ts.get("loai_tai_san"))
            add_field("ghi_chu", fields.get("noi_dung_dang_ky"))

        elif doc_id == "B5":
            add_field("so_bien_ban", fields.get("so_bien_ban"))
            add_field("ngay_lap", fields.get("ngay_lap"))
            
            bg = fields.get("ben_giao", [])
            if bg:
                add_field("ben_giao", bg[0].get("ho_ten"))
            
            bn = fields.get("ben_nhan", {})
            add_field("ben_nhan_ten", bn.get("ten"))
            add_field("ben_nhan_don_vi", bn.get("don_vi"))
            add_field("ben_nhan_nguoi_dai_dien", bn.get("nguoi_dai_dien"))
            add_field("ben_nhan_chuc_danh", bn.get("chuc_danh"))
            add_field("nguoi_chung_kien", fields.get("nguoi_chung_kien"))
            add_field("tong_so_loai_luu_giu", fields.get("tong_so_loai_luu_giu"))
            add_field("tong_so_ban_luu_giu", fields.get("tong_so_ban_luu_giu"))

            bh = fields.get("hop_dong_bao_hiem", {})
            add_field("hop_dong_bao_hiem_so", bh.get("so"))
            add_field("hop_dong_bao_hiem_ngay", bh.get("ngay"))
            add_field("nguoi_thu_huong", bh.get("nguoi_thu_huong"))

            pl = fields.get("phu_luc_doi_chieu", {})
            add_field("tong_chi_tieu", pl.get("tong_chi_tieu"))
            add_field("so_chi_tieu_khop", pl.get("so_chi_tieu_khop"))
            add_field("ket_luan", pl.get("ket_luan"))

            nk = fields.get("nhap_kho", {})
            add_field("ma_ho_so", nk.get("ma_ho_so"))
            add_field("ngay_nhap_kho", nk.get("ngay_nhap_kho"))
            add_field("gio_nhap_kho", nk.get("gio_nhap_kho"))
            add_field("vi_tri_luu_tru", nk.get("vi_tri_luu_tru"))
            add_field("can_bo_nhap_kho", nk.get("can_bo_nhap_kho"))
            add_field("kiem_soat_vien", nk.get("kiem_soat_vien"))

            # bảng danh_muc_ban_giao
            table_data.setdefault("table_danh_muc_ban_giao", []).extend([
                (doc_id, item.get("stt"), item.get("ten"), item.get("loai_ban"), item.get("so_luong"))
                for item in fields.get("danh_muc_ban_giao", [])
            ])

    # Ghi các sheet bảng
    for tbl_name, rows in table_data.items():
        ws_t = wb.create_sheet(tbl_name)
        if "ben_the_chap" in tbl_name:
            t_headers = ["doc_id", "row", "ho_ten", "ngay_sinh", "so_cccd", "ngay_cap_cccd", "noi_thuong_tru"]
        else:
            t_headers = ["doc_id", "row", "ten", "loai_ban", "so_luong"]
        ws_t.append(t_headers)
        for c in ws_t[1]:
            c.fill = fill_header
            c.font = font_header
        for r in rows:
            ws_t.append(list(r))

    # Tự động canh độ rộng cột
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

    wb.save(OUTPUT_PATH)
    print(f"✅ Đã tạo thành công file Ground Truth: {OUTPUT_PATH}")

if __name__ == "__main__":
    main()
