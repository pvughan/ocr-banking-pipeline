"""
eval_metrics.py
===============
Mục 6.2 đề bài – đánh giá chất lượng so với Ground Truth (GT).

CHỈ SỐ
    Field Accuracy (FA) = số trường khớp GT (sau chuẩn hoá) / tổng trường GT
                          -> theo từng tài liệu, từng loại B1..B6, macro trung bình, TÁCH RIÊNG từng
                             biến thể (mobile/scan/scan_bw150) và từng nhóm trường (số tiền, ngày,
                             họ tên, địa chỉ, mã số) – đúng yêu cầu "không gộp chung" của file Lưu ý.
    CER = (S + D + I) / N        edit distance ký tự / độ dài GT   (trường text dài, mặc định >= 20 ký tự)
    WER = edit distance theo TỪ / số từ GT
    Table score = số ô đúng vị trí (hàng, cột) / tổng ô GT   (TEDS chưa cài đặt – ghi chú trong báo cáo)
    Robustness  = chênh lệch FA giữa các biến thể của cùng tài liệu (B1, B2, B5) + tỉ lệ trường nhất quán
    Latency     = trung vị / P95 giây trên mỗi tài liệu (đọc từ sheet metadata của file kết quả)
    Điểm trích xuất = 0.6 × macro-FA + 0.2 × table-score + 0.2 × (1 − CER trung bình)

ĐỊNH DẠNG GT (Excel)
    Sheet `ground_truth` : doc_id | field_key | value          (doc_id = B1..B6; biến thể dùng chung GT)
    Sheet `table_<key>`  : doc_id | row | <column_key>...        (mỗi dòng là một hàng của bảng)
    LƯU Ý: định dạng ô GT là Text để không mất số 0 đầu ("05", "079085...").

CÁCH DÙNG
    python eval_metrics.py --gt GT.xlsx --pred outputs/ --templates templates/ --out metrics.xlsx
    python eval_metrics.py --make-gt-skeleton --template templates/B2_PhieuYeuCau.xlsx --doc-id B2 --out GT_B2.xlsx
Tên file kết quả cần bắt đầu bằng mã tài liệu và chứa biến thể, ví dụ: B1_scan.xlsx, B1_mobile.xlsx,
B1_scan_bw150.xlsx, B2_scan_RESULT.xlsx.
"""
from __future__ import annotations

import argparse
import re
import sys
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    from rapidfuzz.distance import Levenshtein as _Lev
except ImportError:  # pragma: no cover - rapidfuzz nằm trong requirements, đây chỉ là dự phòng
    _Lev = None

from template_parser import FieldSpec, Template, clean_whitespace, comparison_key, parse_template

LONG_TEXT_MIN_CHARS = 20
VARIANTS = ("scan_bw150", "mobile", "scan")
GROUPS = ("so_tien", "ngay_thang", "ho_ten", "dia_chi", "ma_so", "khac")


# --------------------------------------------------------------------------- #
# Edit distance (Levenshtein) – CER / WER
# --------------------------------------------------------------------------- #
def _levenshtein_ops(a: Any, b: Any) -> Tuple[int, int, int]:
    """Trả về (S, D, I): số phép thay / xoá / chèn để biến chuỗi/chuỗi-từ `a` (GT) thành `b` (dự đoán)."""
    if _Lev is not None:
        ops = _Lev.editops(a, b).as_list()
        s = sum(1 for t, _, _ in ops if t == "replace")
        d = sum(1 for t, _, _ in ops if t == "delete")
        i = sum(1 for t, _, _ in ops if t == "insert")
        return s, d, i
    # DP dự phòng (không tách được S/D/I -> quy toàn bộ vào S)
    prev = list(range(len(b) + 1))
    for i, ca in enumerate(a, 1):
        cur = [i]
        for j, cb in enumerate(b, 1):
            cur.append(min(prev[j] + 1, cur[j - 1] + 1, prev[j - 1] + (ca != cb)))
        prev = cur
    return prev[-1], 0, 0


def cer(gt: str, pred: str) -> Tuple[float, int]:
    """CER = (S+D+I)/N với N = độ dài GT. Trả về (cer, số phép sửa)."""
    if not gt:
        return (0.0 if not pred else 1.0), len(pred)
    s, d, i = _levenshtein_ops(gt, pred)
    return (s + d + i) / len(gt), s + d + i


def wer(gt: str, pred: str) -> Tuple[float, int]:
    g, p = gt.split(), pred.split()
    if not g:
        return (0.0 if not p else 1.0), len(p)
    s, d, i = _levenshtein_ops(g, p)
    return (s + d + i) / len(g), s + d + i


# --------------------------------------------------------------------------- #
# Đọc GT / kết quả dự đoán
# --------------------------------------------------------------------------- #
@dataclass
class Prediction:
    path: Path
    doc_id: str
    variant: str
    values: Dict[str, str]
    tables: Dict[str, List[Dict[str, str]]]
    latency_s: Optional[float]
    pages: Optional[int]


def _read_sheet(path: Path, sheet: str) -> pd.DataFrame:
    return pd.read_excel(path, sheet_name=sheet, dtype=str, keep_default_na=False, engine="openpyxl")


def load_ground_truth(path: Path) -> Tuple[Dict[str, Dict[str, str]], Dict[str, Dict[str, List[Dict[str, str]]]]]:
    """-> (values[doc_id][field_key], tables[doc_id][field_key] = list hàng)."""
    wb = load_workbook(path, read_only=True)
    names = wb.sheetnames
    wb.close()
    if "ground_truth" not in names:
        raise ValueError("File GT cần có sheet 'ground_truth' (cột: doc_id, field_key, value).")
    df = _read_sheet(path, "ground_truth")
    for col in ("field_key", "value"):
        if col not in df.columns:
            raise ValueError(f"Sheet ground_truth thiếu cột '{col}'.")
    if "doc_id" not in df.columns:
        df["doc_id"] = "ALL"
    values: Dict[str, Dict[str, str]] = defaultdict(dict)
    for _, r in df.iterrows():
        if r["field_key"].strip():
            values[r["doc_id"].strip() or "ALL"][r["field_key"].strip()] = str(r["value"])
    tables: Dict[str, Dict[str, List[Dict[str, str]]]] = defaultdict(lambda: defaultdict(list))
    for name in names:
        if not name.lower().startswith("table_"):
            continue
        key = name[len("table_"):]
        tdf = _read_sheet(path, name)
        if "doc_id" not in tdf.columns:
            tdf["doc_id"] = "ALL"
        cols = [c for c in tdf.columns if c not in ("doc_id", "row")]
        for _, r in tdf.iterrows():
            tables[r["doc_id"].strip() or "ALL"][key].append({c: str(r[c]) for c in cols})
    return dict(values), {d: dict(t) for d, t in tables.items()}


def parse_doc_and_variant(path: Path) -> Tuple[str, str]:
    stem = path.stem
    m = re.match(r"(?i)^(B\d+)", stem)
    doc_id = m.group(1).upper() if m else "ALL"
    low = stem.lower()
    variant = "unknown"
    for v in VARIANTS:                       # bw150 phải được kiểm tra trước 'scan'
        if v in low or (v == "scan_bw150" and "bw150" in low):
            variant = v
            break
    return doc_id, variant


def load_prediction(path: Path) -> Prediction:
    wb = load_workbook(path, read_only=True, data_only=True)
    names = wb.sheetnames
    wb.close()
    if "Ket_qua" not in names:
        raise ValueError(f"{path.name}: thiếu sheet 'Ket_qua'.")
    kq = _read_sheet(path, "Ket_qua")
    values = {str(r["field_key"]): str(r["value"]) for _, r in kq.iterrows()}
    tables: Dict[str, List[Dict[str, str]]] = {}
    for name in names:
        if name.lower().startswith("table_"):
            tdf = _read_sheet(path, name)
            tables[name[len("table_"):]] = [{c: str(r[c]) for c in tdf.columns if c != "row"}
                                            for _, r in tdf.iterrows()]
    latency, pages = None, None
    if "metadata" in names:
        md = _read_sheet(path, "metadata")
        kv = {str(r.iloc[0]): str(r.iloc[1]) for _, r in md.iterrows()}
        try:
            latency = float(kv.get("Total latency (s)", ""))
        except ValueError:
            pass
        try:
            pages = int(float(kv.get("pages", "")))
        except ValueError:
            pass
    doc_id, variant = parse_doc_and_variant(path)
    return Prediction(path, doc_id, variant, values, tables, latency, pages)


def load_templates(src: Optional[str]) -> Dict[str, Template]:
    """`--templates` là file (dùng cho mọi doc) hoặc thư mục (file có tên bắt đầu bằng mã tài liệu)."""
    out: Dict[str, Template] = {}
    if not src:
        return out
    p = Path(src)
    files = [p] if p.is_file() else sorted(p.glob("*.xlsx"))
    for f in files:
        try:
            tpl = parse_template(f)
        except Exception as exc:
            print(f"[cảnh báo] bỏ qua template {f.name}: {exc}", file=sys.stderr)
            continue
        m = re.match(r"(?i)^(B\d+)", f.stem)
        out[m.group(1).upper() if (m and p.is_dir()) else "*"] = tpl
    return out


# --------------------------------------------------------------------------- #
# Kiểu trường, nhóm trường
# --------------------------------------------------------------------------- #
def infer_type(value: str) -> str:
    v = clean_whitespace(value)
    if re.fullmatch(r"\d{1,2}\s*[/\-.]\s*\d{1,2}\s*[/\-.]\s*\d{4}", v):
        return "date"
    if re.fullmatch(r"\d{1,3}(?:[.,]\d{3}){1,}(?:\s*(?:đ|đồng|vnd))?", v, flags=re.I):
        return "money"
    return "text"


def field_group(key: str, ftype: str, name: str = "") -> str:
    k, n = key.lower(), name.lower()
    if "dia_chi" in k or "địa chỉ" in n or "noi_thuong_tru" in k or "dia_diem" in k:
        return "dia_chi"
    if ftype == "money":
        return "so_tien"
    if ftype == "date":
        return "ngay_thang"
    if "ho_ten" in k or k.endswith("_ten") or k.startswith("ten_") or "nguoi_" in k or "cong_chung_vien" in k:
        return "ho_ten"
    if ftype == "number" or re.search(r"(^|[._])(so|ma|msdn|mst|cccd)", k):
        return "ma_so"
    return "khac"


def _spec_for(templates: Dict[str, Template], doc_id: str, key: str) -> Optional[FieldSpec]:
    tpl = templates.get(doc_id) or templates.get("*")
    return tpl.get(key) if tpl else None


# --------------------------------------------------------------------------- #
# Đánh giá một cặp (GT, dự đoán)
# --------------------------------------------------------------------------- #
def evaluate_prediction(pred: Prediction, gt_values: Dict[str, str], gt_tables: Dict[str, List[Dict[str, str]]],
                        templates: Dict[str, Template]) -> Dict[str, Any]:
    fields: List[Dict[str, Any]] = []
    for key, gt_raw in gt_values.items():
        if not clean_whitespace(gt_raw):
            continue                                  # GT rỗng: không tính (không có "đáp án đúng")
        spec = _spec_for(templates, pred.doc_id, key)
        ftype = spec.type if spec else infer_type(gt_raw)
        enums = spec.enum_values if spec else None
        pred_raw = pred.values.get(key, "")
        k_gt = comparison_key(ftype, gt_raw, enums)
        k_pr = comparison_key(ftype, pred_raw, enums)
        row: Dict[str, Any] = {
            "doc_id": pred.doc_id, "variant": pred.variant, "file": pred.path.name, "field_key": key,
            "type": ftype, "group": field_group(key, ftype, spec.field_name if spec else ""),
            "gt": gt_raw, "pred": pred_raw, "match": k_gt == k_pr and bool(k_gt), "missing": not pred_raw.strip(),
        }
        if ftype == "text":                         # CER/WER trên dạng đã chuẩn hoá (casefold, gọn khoảng trắng)
            row["cer"], row["cer_edits"] = cer(k_gt, k_pr)
            row["wer"], row["wer_edits"] = wer(k_gt, k_pr)
            row["gt_len"] = len(k_gt)
        fields.append(row)

    # bảng: so từng ô theo vị trí (hàng, cột)
    table_rows: List[Dict[str, Any]] = []
    for key, gt_rows in gt_tables.items():
        spec = _spec_for(templates, pred.doc_id, key)
        col_types = {c.column_key: c.type for c in spec.columns} if spec else {}
        pr_rows = pred.tables.get(key, [])
        total = correct = 0
        for i, g_row in enumerate(gt_rows):
            p_row = pr_rows[i] if i < len(pr_rows) else {}
            for col, g_val in g_row.items():
                if not clean_whitespace(g_val):
                    continue
                total += 1
                ctype = col_types.get(col, infer_type(g_val))
                if comparison_key(ctype, g_val) == comparison_key(ctype, p_row.get(col, "")) and g_val.strip():
                    correct += 1
        table_rows.append({"doc_id": pred.doc_id, "variant": pred.variant, "file": pred.path.name,
                           "table": key, "gt_rows": len(gt_rows), "pred_rows": len(pr_rows),
                           "cells": total, "cells_correct": correct,
                           "table_score": correct / total if total else np.nan})
    return {"fields": fields, "tables": table_rows}


# --------------------------------------------------------------------------- #
# Tổng hợp
# --------------------------------------------------------------------------- #
def _fa(df: pd.DataFrame) -> float:
    return float(df["match"].mean()) if len(df) else float("nan")


def evaluate(gt_path: Path, pred_paths: List[Path], templates: Dict[str, Template]) -> Dict[str, pd.DataFrame]:
    gt_values, gt_tables = load_ground_truth(gt_path)
    all_fields: List[Dict[str, Any]] = []
    all_tables: List[Dict[str, Any]] = []
    meta_rows: List[Dict[str, Any]] = []
    skipped: List[str] = []
    for p in pred_paths:
        try:
            pred = load_prediction(p)
        except Exception as exc:
            skipped.append(f"{p.name}: {exc}")
            continue
        gv = gt_values.get(pred.doc_id) or gt_values.get("ALL")
        if gv is None:
            skipped.append(f"{p.name}: GT không có doc_id '{pred.doc_id}'")
            continue
        gtab = gt_tables.get(pred.doc_id) or gt_tables.get("ALL") or {}
        res = evaluate_prediction(pred, gv, gtab, templates)
        all_fields += res["fields"]
        all_tables += res["tables"]
        meta_rows.append({"doc_id": pred.doc_id, "variant": pred.variant, "file": p.name,
                          "latency_s": pred.latency_s, "pages": pred.pages})
    for s in skipped:
        print(f"[cảnh báo] bỏ qua {s}", file=sys.stderr)
    if not all_fields:
        raise SystemExit("Không có file kết quả nào khớp GT để đánh giá.")

    fields = pd.DataFrame(all_fields)
    tables = pd.DataFrame(all_tables)
    meta = pd.DataFrame(meta_rows)

    # --- theo từng file (tài liệu × biến thể)
    per_doc_rows = []
    for (doc, var, f), g in fields.groupby(["doc_id", "variant", "file"], sort=True):
        text = g[g["type"] == "text"].dropna(subset=["cer"]) if "cer" in g else g.iloc[0:0]
        long_text = text[text["gt_len"] >= LONG_TEXT_MIN_CHARS]
        tsub = tables[tables["file"] == f] if len(tables) else tables
        m = meta[meta["file"] == f].iloc[0]
        per_doc_rows.append({
            "doc_id": doc, "variant": var, "file": f, "n_fields": len(g), "n_correct": int(g["match"].sum()),
            "FA": _fa(g), "n_missing": int(g["missing"].sum()),
            "CER_long": float(long_text["cer"].mean()) if len(long_text) else np.nan,
            "CER_text": float(text["cer"].mean()) if len(text) else np.nan,
            "WER": float(long_text["wer"].mean()) if len(long_text) else
                   (float(text["wer"].mean()) if len(text) else np.nan),
            "table_score": float(tsub["table_score"].mean()) if len(tsub) else np.nan,
            "latency_s": m["latency_s"], "pages": m["pages"]})
    per_doc = pd.DataFrame(per_doc_rows)

    # --- theo loại tài liệu (macro trên các biến thể)
    per_type = per_doc.groupby("doc_id").agg(FA=("FA", "mean"), CER_long=("CER_long", "mean"),
                                             table_score=("table_score", "mean"),
                                             n_files=("file", "count")).reset_index()
    # --- theo biến thể (KHÔNG gộp chung)
    per_variant = per_doc.groupby("variant").agg(FA=("FA", "mean"), CER_long=("CER_long", "mean"),
                                                 table_score=("table_score", "mean"),
                                                 n_files=("file", "count")).reset_index()
    # --- theo nhóm trường × biến thể
    grp = fields.groupby(["variant", "group"]).agg(FA=("match", "mean"), n_fields=("match", "size")).reset_index()
    per_group = grp.pivot(index="group", columns="variant", values="FA").reindex(
        [g for g in GROUPS if g in grp["group"].unique()])
    per_group["ALL_variants"] = fields.groupby("group")["match"].mean()
    per_group = per_group.reset_index()

    # --- robustness
    rob_rows = []
    for doc, g in per_doc.groupby("doc_id"):
        if g["variant"].nunique() < 2:
            continue
        by_var = g.groupby("variant")["FA"].mean()
        sub = fields[fields["doc_id"] == doc]
        same = []
        for _, gk in sub.groupby("field_key"):       # trường có cho cùng kết quả ở mọi biến thể không?
            same.append(len({comparison_key(t, v) for t, v in zip(gk["type"], gk["pred"])}) == 1)
        consistent = float(np.mean(same)) if same else np.nan
        rob_rows.append({"doc_id": doc, **{f"FA_{v}": by_var.get(v, np.nan) for v in by_var.index},
                         "FA_spread(max-min)": float(by_var.max() - by_var.min()),
                         "field_consistency": consistent})
    robustness = pd.DataFrame(rob_rows)

    # --- latency
    lat = per_doc["latency_s"].dropna()
    latency = pd.DataFrame([{"n_docs": len(lat), "median_s": float(lat.median()) if len(lat) else np.nan,
                             "p95_s": float(np.percentile(lat, 95)) if len(lat) else np.nan,
                             "median_s_per_page": float((per_doc["latency_s"] / per_doc["pages"]).median())
                             if per_doc["pages"].notna().any() else np.nan}])

    # --- điểm trích xuất: 0.6*macroFA + 0.2*table + 0.2*(1-CER)
    macro_fa = float(per_type["FA"].mean())
    table_macro = float(per_type["table_score"].mean()) if per_type["table_score"].notna().any() else np.nan
    cer_avg = float(per_type["CER_long"].mean()) if per_type["CER_long"].notna().any() else \
        float(per_doc["CER_text"].mean()) if per_doc["CER_text"].notna().any() else np.nan
    weights = {"macro_FA": (0.6, macro_fa), "table_score": (0.2, table_macro),
               "1-CER": (0.2, (1 - min(max(cer_avg, 0.0), 1.0)) if not np.isnan(cer_avg) else np.nan)}
    avail = {k: v for k, v in weights.items() if not np.isnan(v[1])}
    wsum = sum(w for w, _ in avail.values())
    score = sum(w * v for w, v in avail.values()) / wsum if wsum else np.nan   # chuẩn hoá nếu thiếu thành phần
    summary = pd.DataFrame([{
        "macro_FA": macro_fa, "table_score": table_macro, "CER_avg": cer_avg,
        "WER_avg": float(per_doc["WER"].mean()) if per_doc["WER"].notna().any() else np.nan,
        "extraction_score(0.6FA+0.2TBL+0.2(1-CER))": score,
        "note": "" if len(avail) == 3 else "thiếu thành phần (không có bảng/trường text dài) -> trọng số được chuẩn hoá lại"}])

    return {"summary": summary, "per_type": per_type, "per_variant": per_variant, "per_group": per_group,
            "robustness": robustness, "latency": latency, "per_document": per_doc,
            "field_details": fields, "table_details": tables}


# --------------------------------------------------------------------------- #
# Xuất báo cáo
# --------------------------------------------------------------------------- #
def print_report(tables: Dict[str, pd.DataFrame]) -> None:
    pd.set_option("display.width", 200)
    pd.set_option("display.max_columns", 30)
    fmt = lambda x: f"{x:.3f}" if isinstance(x, float) else x
    for name in ("summary", "per_type", "per_variant", "per_group", "robustness", "latency", "per_document"):
        df = tables.get(name)
        if df is None or df.empty:
            continue
        print(f"\n=== {name.upper()} ===")
        print(df.to_string(index=False, float_format=lambda x: f"{x:.3f}"))


def write_report_excel(tables: Dict[str, pd.DataFrame], path: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for name, df in tables.items():
        ws = wb.create_sheet(name[:31])
        ws.append(list(df.columns))
        for row in df.itertuples(index=False):
            ws.append([None if (isinstance(v, float) and np.isnan(v)) else (bool(v) if isinstance(v, np.bool_) else v)
                       for v in row])
        for c in ws[1]:
            c.fill = PatternFill("solid", start_color="4F46E5", end_color="4F46E5")
            c.font = Font(bold=True, color="FFFFFF")
            c.alignment = Alignment(horizontal="center", wrap_text=True)
        for j, col in enumerate(ws.columns, start=1):
            longest = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[get_column_letter(j)].width = min(60, max(10, longest + 2))
            if isinstance(ws.cell(row=2, column=j).value, float):
                for c in col[1:]:
                    c.number_format = "0.000"
        ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def make_gt_skeleton(template: Template, doc_id: str, out: Path) -> None:
    """Sinh khung file GT từ template để nhóm điền tay (khuyến nghị 2 người đọc chéo – đề bài mục 13)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "ground_truth"
    ws.append(["doc_id", "field_key", "value"])
    for f in template.fields:
        if not f.is_table:
            ws.append([doc_id, f.field_key, ""])
    for c in ws[1]:
        c.fill = PatternFill("solid", start_color="4F46E5", end_color="4F46E5")
        c.font = Font(bold=True, color="FFFFFF")
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        for c in row:
            c.number_format = "@"                 # Text: giữ nguyên số 0 đầu
    for f in template.table_fields:
        wt = wb.create_sheet(f"table_{f.field_key}"[:31])
        wt.append(["doc_id", "row"] + [c.column_key for c in f.columns])
        for c in wt[1]:
            c.fill = PatternFill("solid", start_color="4F46E5", end_color="4F46E5")
            c.font = Font(bold=True, color="FFFFFF")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


# --------------------------------------------------------------------------- #
def main(argv: Optional[List[str]] = None) -> int:
    ap = argparse.ArgumentParser(description="Đánh giá OCR: FA / CER / WER / table / robustness / latency")
    ap.add_argument("--gt", help="file Excel Ground Truth")
    ap.add_argument("--pred", help="thư mục chứa các file kết quả .xlsx (hoặc một file)")
    ap.add_argument("--templates", help="template (file hoặc thư mục) để biết type từng trường")
    ap.add_argument("--out", default="outputs/metrics.xlsx", help="file Excel báo cáo")
    ap.add_argument("--make-gt-skeleton", action="store_true", help="sinh khung GT từ --template")
    ap.add_argument("--template", help="template dùng cho --make-gt-skeleton")
    ap.add_argument("--doc-id", default="B2")
    args = ap.parse_args(argv)

    if args.make_gt_skeleton:
        if not args.template:
            ap.error("--make-gt-skeleton cần --template")
        make_gt_skeleton(parse_template(args.template), args.doc_id, Path(args.out))
        print(f"Đã tạo khung GT: {args.out}")
        return 0
    if not args.gt or not args.pred:
        ap.error("cần --gt và --pred")
    p = Path(args.pred)
    preds = [p] if p.is_file() else sorted(x for x in p.glob("*.xlsx") if not x.name.startswith(("~", "metrics")))
    if not preds:
        ap.error(f"không tìm thấy file .xlsx trong {p}")
    tables = evaluate(Path(args.gt), preds, load_templates(args.templates))
    print_report(tables)
    write_report_excel(tables, Path(args.out))
    print(f"\nĐã ghi báo cáo: {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
