"""
template_parser.py
==================
Mục 4.1 của đề bài: đọc Template Excel + Validators + Post-processing.

Nội dung
--------
1. Mô hình dữ liệu: ``ColumnSpec``, ``FieldSpec``, ``Template``.
2. ``parse_template``  : đọc template Excel (8 cột chuẩn), hỗ trợ trường ``table`` với cột con
   khai báo trong ``notes`` (``columns=a|b|c``) HOẶC sheet phụ ``table_<field_key>``.
3. ``template_to_excel_bytes`` / ``template_from_dataframe``: tạo & sửa template trong app, export
   lại đúng định dạng chuẩn.
4. Validators: regex, ``enum:a|b|c``, ``number``, ``msdn10``, ``cccd12``, ``email``,
   ``check_digit:cccd``, ``check_digit:mst``, ``range:a-b``...
5. Post-processing: ngày -> dd/MM/yyyy, tiền -> số nguyên, proper-case tên riêng, sửa dấu
   tiếng Việt thường gặp, checkbox -> enum/boolean, số bằng chữ -> số.
6. ``comparison_key``: khoá so sánh dùng chung cho Fusion và Eval (tránh "bẫy" normalize sai
   trước khi so sánh hai engine – đề bài mục 13).
"""
from __future__ import annotations

import datetime as dt
import io
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, BinaryIO, Dict, Iterable, List, Optional, Sequence, Tuple, Union

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from rapidfuzz import fuzz

# --------------------------------------------------------------------------- #
# Hằng số
# --------------------------------------------------------------------------- #
VALID_TYPES = ("text", "number", "date", "money", "boolean", "checkbox", "table")
TEMPLATE_COLUMNS = ["field_key", "field_name", "type", "required", "validation", "anchor", "page", "notes"]
TABLE_SHEET_COLUMNS = ["column_key", "column_name", "type", "validation", "notes"]
INDIGO = "4F46E5"


class TemplateError(Exception):
    """Template không hợp lệ (thiếu cột bắt buộc, trùng field_key...)."""


# --------------------------------------------------------------------------- #
# Tiện ích văn bản tiếng Việt
# --------------------------------------------------------------------------- #
def nfc(text: str) -> str:
    return unicodedata.normalize("NFC", text or "")


def strip_diacritics(text: str) -> str:
    """Bỏ dấu tiếng Việt (đ -> d)."""
    text = unicodedata.normalize("NFD", text or "")
    text = "".join(c for c in text if unicodedata.category(c) != "Mn")
    return text.replace("đ", "d").replace("Đ", "D")


def fold_vietnamese(text: str) -> str:
    """lowercase + bỏ dấu + chỉ giữ chữ/số, gộp khoảng trắng. Dùng để dò anchor/so khớp mờ."""
    t = strip_diacritics(nfc(text)).lower()
    t = re.sub(r"[^a-z0-9]+", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def clean_whitespace(text: str) -> str:
    text = nfc(text).replace("\u00a0", " ").replace("\u200b", "")
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r"\s*\n\s*", " ", text)
    return text.strip()


def _s(value: Any) -> str:
    """Ép về str an toàn (None/NaN -> '')."""
    if value is None:
        return ""
    if isinstance(value, float) and value != value:  # NaN
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


# --------------------------------------------------------------------------- #
# Mô hình dữ liệu
# --------------------------------------------------------------------------- #
@dataclass
class ColumnSpec:
    column_key: str
    column_name: str = ""
    type: str = "text"
    validation: str = ""
    notes: str = ""

    @property
    def key(self) -> str:
        return self.column_key

    @property
    def required(self) -> bool:
        return False

    @property
    def wants_proper_case(self) -> bool:
        return "proper" in self.notes.lower()

    @property
    def enum_values(self) -> List[str]:
        return _enum_values(self.validation)


@dataclass
class FieldSpec:
    field_key: str
    field_name: str = ""
    type: str = "text"
    required: bool = False
    validation: str = ""
    anchor: str = ""
    page: str = "all"
    notes: str = ""
    columns: List[ColumnSpec] = field(default_factory=list)

    @property
    def key(self) -> str:
        return self.field_key

    @property
    def is_table(self) -> bool:
        return self.type == "table"

    @property
    def wants_proper_case(self) -> bool:
        return "proper" in self.notes.lower()

    @property
    def enum_values(self) -> List[str]:
        return _enum_values(self.validation)

    def page_candidates(self, n_pages: int) -> List[int]:
        """
        Trang dự kiến (1-based) hợp lệ với tài liệu `n_pages` trang. Đây chỉ là GỢI Ý (đề bài:
        "pipeline không được phụ thuộc cứng"): không parse được -> trả về tất cả các trang.
        """
        spec = (self.page or "").strip().lower()
        everything = list(range(1, n_pages + 1))
        if spec in ("", "all", "*", "any"):
            return everything
        pages: List[int] = []
        for part in re.split(r"[;,]", spec):
            part = part.strip()
            m = re.fullmatch(r"(\d+)\s*-\s*(\d+)", part)
            if m:
                a, b = int(m.group(1)), int(m.group(2))
                pages.extend(range(min(a, b), max(a, b) + 1))
            elif part.isdigit():
                pages.append(int(part))
        pages = sorted({p for p in pages if 1 <= p <= n_pages})
        return pages or everything


@dataclass
class Template:
    name: str
    fields: List[FieldSpec]
    source: str = ""
    warnings: List[str] = field(default_factory=list)

    @property
    def keys(self) -> List[str]:
        return [f.field_key for f in self.fields]

    @property
    def table_fields(self) -> List[FieldSpec]:
        return [f for f in self.fields if f.is_table]

    def get(self, key: str) -> Optional[FieldSpec]:
        for f in self.fields:
            if f.field_key == key:
                return f
        return None

    def to_dataframe(self) -> pd.DataFrame:
        rows = []
        for f in self.fields:
            notes = f.notes
            if f.is_table and f.columns and "columns=" not in notes.lower():
                cols = "|".join(c.column_key for c in f.columns)
                notes = (notes + "; " if notes else "") + f"columns={cols}"
            rows.append({
                "field_key": f.field_key, "field_name": f.field_name, "type": f.type,
                "required": "Y" if f.required else "N", "validation": f.validation,
                "anchor": f.anchor, "page": f.page, "notes": notes,
            })
        return pd.DataFrame(rows, columns=TEMPLATE_COLUMNS)


# --------------------------------------------------------------------------- #
# Đọc template
# --------------------------------------------------------------------------- #
def _enum_values(validation: str) -> List[str]:
    v = (validation or "").strip()
    if v.lower().startswith("enum:"):
        return [x.strip() for x in v[5:].split("|") if x.strip()]
    return []


def _open_workbook(source: Union[str, Path, bytes, BinaryIO, Any]):
    if isinstance(source, (str, Path)):
        return load_workbook(str(source), data_only=True)
    if isinstance(source, (bytes, bytearray)):
        return load_workbook(io.BytesIO(bytes(source)), data_only=True)
    if hasattr(source, "getvalue"):  # Streamlit UploadedFile / BytesIO
        return load_workbook(io.BytesIO(source.getvalue()), data_only=True)
    if hasattr(source, "read"):
        return load_workbook(io.BytesIO(source.read()), data_only=True)
    raise TemplateError(f"Kiểu nguồn template không hỗ trợ: {type(source)!r}")


def _norm_header(h: Any) -> str:
    return re.sub(r"\s+", "_", _s(h).lower())


def _read_sheet_table(ws, expected: Sequence[str], min_hits: int = 3) -> Tuple[List[str], List[List[Any]]]:
    """Tìm dòng header (>= min_hits cột chuẩn) rồi trả về (header, các dòng dữ liệu)."""
    rows = list(ws.iter_rows(values_only=True))
    for i, row in enumerate(rows[:15]):
        headers = [_norm_header(c) for c in row]
        if sum(1 for h in headers if h in expected) >= min_hits:
            data = [list(r) for r in rows[i + 1:] if any(_s(c) for c in r)]
            return headers, data
    raise TemplateError(f"Sheet '{ws.title}' không có dòng tiêu đề chuẩn {list(expected)}")


def _parse_bool(v: Any) -> bool:
    return _s(v).lower() in ("y", "yes", "true", "1", "có", "co", "x")


def _columns_from_notes(notes: str) -> List[ColumnSpec]:
    m = re.search(r"columns\s*=\s*([^\n;]+)", notes or "", flags=re.I)
    if not m:
        return []
    cols: List[ColumnSpec] = []
    for token in m.group(1).split("|"):
        token = token.strip()
        if not token:
            continue
        key, _, ctype = token.partition(":")
        key = key.strip()
        ctype = ctype.strip().lower() or "text"
        if ctype not in VALID_TYPES or ctype == "table":
            ctype = "text"
        cols.append(ColumnSpec(column_key=key, column_name=key.replace("_", " ").title(), type=ctype))
    return cols


def _columns_from_sheet(wb, field_key: str) -> List[ColumnSpec]:
    wanted = f"table_{field_key}".lower()
    for ws in wb.worksheets:
        if ws.title.strip().lower() == wanted:
            headers, data = _read_sheet_table(ws, TABLE_SHEET_COLUMNS, min_hits=2)
            idx = {h: i for i, h in enumerate(headers)}
            cols = []
            for r in data:
                get = lambda name: _s(r[idx[name]]) if name in idx and idx[name] < len(r) else ""
                key = get("column_key")
                if not key:
                    continue
                ctype = get("type").lower() or "text"
                cols.append(ColumnSpec(
                    column_key=key, column_name=get("column_name") or key,
                    type=ctype if ctype in VALID_TYPES and ctype != "table" else "text",
                    validation=get("validation"), notes=get("notes")))
            return cols
    return []


def parse_template(source: Union[str, Path, bytes, BinaryIO, Any], name: Optional[str] = None) -> Template:
    """
    Đọc template Excel chuẩn mục 4.1.

    Ưu tiên cột con của bảng từ sheet phụ ``table_<field_key>``; nếu không có thì đọc
    ``columns=a|b|c`` trong cột notes. Trả về ``Template`` kèm danh sách cảnh báo (không ném lỗi
    cho các vấn đề nhẹ như type lạ) – lỗi nặng (thiếu cột, trùng field_key) ném ``TemplateError``.
    """
    try:
        wb = _open_workbook(source)
    except TemplateError:
        raise
    except Exception as exc:  # file hỏng, không phải xlsx...
        raise TemplateError(f"Không đọc được file Excel template: {exc}") from exc

    ws = next((w for w in wb.worksheets if w.title.strip().lower() == "template"), None)
    if ws is None:
        ws = next((w for w in wb.worksheets if "field_key" in
                   [_norm_header(c) for r in w.iter_rows(max_row=5, values_only=True) for c in r]), None)
    if ws is None:
        raise TemplateError("Không tìm thấy sheet 'template' (hoặc sheet có cột field_key).")

    headers, data = _read_sheet_table(ws, TEMPLATE_COLUMNS, min_hits=4)
    missing = [c for c in ("field_key", "field_name", "type") if c not in headers]
    if missing:
        raise TemplateError(f"Template thiếu cột bắt buộc: {missing}")
    idx = {h: i for i, h in enumerate(headers) if h}

    def cell(row: List[Any], col: str) -> str:
        i = idx.get(col)
        return _s(row[i]) if i is not None and i < len(row) else ""

    warnings: List[str] = []
    fields: List[FieldSpec] = []
    seen: set = set()
    for n, row in enumerate(data, start=2):
        key = cell(row, "field_key")
        if not key:
            continue
        if key in seen:
            raise TemplateError(f"field_key bị trùng: '{key}' (dòng {n})")
        seen.add(key)
        if not re.fullmatch(r"[A-Za-z0-9_.\[\]]+", key):
            warnings.append(f"field_key '{key}' nên là snake_case không dấu/khoảng trắng.")
        ftype = cell(row, "type").lower() or "text"
        if ftype not in VALID_TYPES:
            warnings.append(f"Trường '{key}': type '{ftype}' không hợp lệ -> dùng 'text'.")
            ftype = "text"
        spec = FieldSpec(
            field_key=key, field_name=cell(row, "field_name") or key, type=ftype,
            required=_parse_bool(cell(row, "required")), validation=cell(row, "validation"),
            anchor=cell(row, "anchor"), page=cell(row, "page") or "all", notes=cell(row, "notes"))
        if ftype == "table":
            spec.columns = _columns_from_sheet(wb, key) or _columns_from_notes(spec.notes)
            if not spec.columns:
                warnings.append(f"Trường table '{key}' chưa khai báo cột con "
                                f"(notes 'columns=a|b|c' hoặc sheet 'table_{key}').")
        if ftype == "checkbox" and not spec.enum_values:
            pass  # checkbox boolean thuần: hợp lệ
        if spec.validation and not _validation_is_known(spec.validation):
            warnings.append(f"Trường '{key}': quy tắc validation '{spec.validation}' không nhận diện được.")
        fields.append(spec)

    if not fields:
        raise TemplateError("Template không có trường nào.")
    if len(fields) < 12:
        warnings.append(f"Template chỉ có {len(fields)} trường (đề bài yêu cầu >= 12 trường).")

    tpl_name = name
    if not tpl_name:
        tpl_name = Path(source).stem if isinstance(source, (str, Path)) else \
            Path(getattr(source, "name", "template")).stem
    return Template(name=tpl_name, fields=fields,
                    source=str(source) if isinstance(source, (str, Path)) else tpl_name,
                    warnings=warnings)


def template_from_dataframe(df: pd.DataFrame, name: str = "template_tuy_chinh") -> Template:
    """Dựng Template từ DataFrame (dùng cho trình soạn template trong app)."""
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        out = df.copy()
        for c in TEMPLATE_COLUMNS:
            if c not in out.columns:
                out[c] = ""
        out[TEMPLATE_COLUMNS].to_excel(xw, sheet_name="template", index=False)
    return parse_template(buf.getvalue(), name=name)


def template_to_excel_bytes(template: Template) -> bytes:
    """Export Template ra Excel đúng định dạng chuẩn (sheet `template` + `table_<key>`)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "template"
    thin = Side(style="thin", color="CBD5E1")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.append(TEMPLATE_COLUMNS)
    for row in template.to_dataframe().itertuples(index=False):
        ws.append(list(row))
    _style_sheet(ws, border)
    for f in template.table_fields:
        title = f"table_{f.field_key}"[:31]
        wt = wb.create_sheet(title)
        wt.append(TABLE_SHEET_COLUMNS)
        for c in f.columns:
            wt.append([c.column_key, c.column_name, c.type, c.validation, c.notes])
        _style_sheet(wt, border)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _style_sheet(ws, border: Border) -> None:
    for c in ws[1]:
        c.fill = PatternFill("solid", start_color=INDIGO)
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center")
    for row in ws.iter_rows():
        for c in row:
            c.border = border
            if c.row > 1:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, col in enumerate(ws.columns, start=1):
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(12, width + 2), 60)
    ws.freeze_panes = "A2"


# --------------------------------------------------------------------------- #
# VALIDATORS
# --------------------------------------------------------------------------- #
@dataclass
class ValidationResult:
    status: str      # ok | warning | error
    message: str = ""

    @property
    def ok(self) -> bool:
        return self.status == "ok"

    @property
    def passed(self) -> bool:
        """Không phải error (ok hoặc warning)."""
        return self.status != "error"


_NAMED_VALIDATORS = {
    "number", "msdn10", "mst10", "cccd12", "cmnd9", "cccd9", "cccd9_12", "email", "phone", "date",
    "check_digit:cccd", "check_digit:mst",
}
DATE_RE = re.compile(r"(?<!\d)(\d{1,2})\s*[/\-.]\s*(\d{1,2})\s*[/\-.]\s*(\d{4})(?!\d)")
DATE_WORDS_RE = re.compile(r"(?:ng[aàáạảãâầấậẩẫăằắặẳẵ]y\s*)?(\d{1,2})\s*th[aáàạảã]ng\s*(\d{1,2})\s*n[aăâ]m\s*(\d{4})",
                           re.I)
MONEY_RE = re.compile(r"(?<!\d)\d{1,3}(?:[.,\s]\d{3})+(?!\d)|(?<!\d)\d{4,}(?!\d)")
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
PHONE_RE = re.compile(r"(?<!\d)0\d{9,10}(?!\d)")


def _validation_is_known(spec: str) -> bool:
    s = (spec or "").strip()
    if not s:
        return True
    low = s.lower()
    if low.startswith(("enum:", "range:", "regex:")) or s.startswith("^"):
        return True
    return all(tok.strip().lower() in _NAMED_VALIDATORS for tok in s.split("|"))


def _mst_checksum_ok(digits10: str) -> bool:
    """Mã số thuế 10 số: C10 = 10 - (sum(Ci*Wi) mod 11) với W = 31,29,23,19,17,13,7,5,3."""
    weights = (31, 29, 23, 19, 17, 13, 7, 5, 3)
    s = sum(int(d) * w for d, w in zip(digits10[:9], weights))
    c = 10 - (s % 11)
    return c < 10 and c == int(digits10[9])


def _check_named(name: str, v: str) -> ValidationResult:
    n = name.strip().lower()
    digits = re.sub(r"[\s.\-]", "", v)
    if n == "number":
        ok = re.fullmatch(r"-?\d+(?:[.,]\d+)?", v.replace(" ", "")) is not None
        return ValidationResult("ok" if ok else "error", "" if ok else "không phải số")
    if n in ("msdn10", "mst10", "check_digit:mst"):
        m = re.fullmatch(r"(\d{10})(?:-?\d{3})?", digits)
        if not m:
            return ValidationResult("error", "MST/MSDN phải có 10 chữ số (hoặc 13 số có đuôi chi nhánh)")
        if not _mst_checksum_ok(m.group(1)):
            return ValidationResult("warning", "MST sai checksum (có thể là dữ liệu giả lập hoặc OCR nhầm số)")
        return ValidationResult("ok")
    if n in ("cccd12", "check_digit:cccd"):
        if not re.fullmatch(r"\d{12}", digits):
            return ValidationResult("error", "CCCD phải gồm đúng 12 chữ số")
        # CCCD không có checksum công khai -> kiểm tra cấu trúc: 3 số đầu = mã tỉnh (001-096).
        if not (1 <= int(digits[:3]) <= 96):
            return ValidationResult("error", "Mã tỉnh/thành (3 số đầu) của CCCD ngoài khoảng 001-096")
        return ValidationResult("ok")
    if n in ("cmnd9", "cccd9"):
        ok = re.fullmatch(r"\d{9}", digits) is not None
        return ValidationResult("ok" if ok else "error", "" if ok else "CMND phải gồm đúng 9 chữ số")
    if n == "cccd9_12":
        ok = re.fullmatch(r"\d{9}|\d{12}", digits) is not None
        return ValidationResult("ok" if ok else "error", "" if ok else "CMND/CCCD phải có 9 hoặc 12 số")
    if n == "email":
        ok = EMAIL_RE.fullmatch(v) is not None
        return ValidationResult("ok" if ok else "error", "" if ok else "email không hợp lệ")
    if n == "phone":
        ok = PHONE_RE.fullmatch(digits) is not None
        return ValidationResult("ok" if ok else "error", "" if ok else "số điện thoại không hợp lệ")
    if n == "date":
        ok = _parse_date(v) is not None
        return ValidationResult("ok" if ok else "error", "" if ok else "ngày không hợp lệ")
    return ValidationResult("ok")


def validate_value(validation: str, value: Optional[str], ftype: str = "text") -> ValidationResult:
    """
    Kiểm tra `value` theo cột `validation` của template.

    Hỗ trợ: regex (bắt đầu bằng ``^`` hoặc ``regex:``), ``enum:a|b|c``, ``range:0-100``,
    validator có tên (``number``, ``msdn10``, ``cccd12``, ``email``, ``check_digit:cccd``...)
    và tổ hợp OR bằng ``|`` (ví dụ ``msdn10|cccd12``). Kiểu ``date`` luôn kiểm tra ngày hợp lệ
    trong lịch. Giá trị rỗng được xử lý ở tầng gọi (missing), ở đây trả về ``ok``.
    """
    spec = (validation or "").strip()
    v = clean_whitespace(value or "")
    if not v:
        return ValidationResult("ok")

    if ftype == "date" and _parse_date(v) is None:
        return ValidationResult("error", "ngày không hợp lệ")

    if not spec:
        return ValidationResult("ok")

    low = spec.lower()
    if low.startswith("enum:"):
        allowed = _enum_values(spec)
        if v in allowed or v.lower() in [a.lower() for a in allowed]:
            return ValidationResult("ok")
        return ValidationResult("error", f"giá trị '{v}' không thuộc enum {allowed}")

    if low.startswith("range:"):
        m = re.fullmatch(r"range:\s*(-?[\d.]+)\s*-\s*(-?[\d.]+)", spec, flags=re.I)
        num = parse_decimal(v)
        if not m or num is None:
            return ValidationResult("error", "không phải số để kiểm tra range")
        lo, hi = Decimal(m.group(1)), Decimal(m.group(2))
        return ValidationResult("ok" if lo <= num <= hi else "error", "" if lo <= num <= hi else
                                f"ngoài khoảng {lo}-{hi}")

    if spec.startswith("^") or low.startswith("regex:"):
        pattern = spec[6:].strip() if low.startswith("regex:") else spec
        try:
            ok = re.search(pattern, v) is not None
        except re.error:
            return ValidationResult("warning", f"regex của template không biên dịch được: {pattern}")
        return ValidationResult("ok" if ok else "error", "" if ok else f"không khớp regex {pattern}")

    tokens = [t.strip() for t in spec.split("|") if t.strip()]
    if tokens and all(t.lower() in _NAMED_VALIDATORS for t in tokens):
        results = [_check_named(t, v) for t in tokens]
        for status in ("ok", "warning"):
            for r in results:
                if r.status == status:
                    return r
        return ValidationResult("error", "; ".join(r.message for r in results if r.message))
    return ValidationResult("ok", "validation không nhận diện được -> bỏ qua")


def validator_search_regex(validation: str, ftype: str = "text") -> Optional[re.Pattern]:
    """Regex dạng KHÔNG neo ^$ để QUÉT trong một cửa sổ văn bản (Tesseract regex-guided search)."""
    spec = (validation or "").strip()
    low = spec.lower()
    parts: List[str] = []
    if ftype == "date":   # luôn dò mọi dạng ngày (kể cả 'ngày .. tháng .. năm'); chuẩn hoá làm sau
        return re.compile(DATE_RE.pattern + "|" + DATE_WORDS_RE.pattern, re.I)
    if spec.startswith("^") or low.startswith("regex:"):
        pattern = spec[6:].strip() if low.startswith("regex:") else spec
        pattern = pattern[1:] if pattern.startswith("^") else pattern
        pattern = pattern[:-1] if pattern.endswith("$") and not pattern.endswith("\\$") else pattern
        try:
            return re.compile(pattern)
        except re.error:
            return None
    for tok in [t.strip().lower() for t in spec.split("|")] if spec and not low.startswith("enum:") else []:
        if tok in ("msdn10", "mst10", "check_digit:mst"):
            parts.append(r"(?<!\d)\d{10}(?:-\d{3})?(?!\d)")
        elif tok in ("cccd12", "check_digit:cccd"):
            parts.append(r"(?<!\d)\d{12}(?!\d)")
        elif tok in ("cmnd9", "cccd9"):
            parts.append(r"(?<!\d)\d{9}(?!\d)")
        elif tok == "cccd9_12":
            parts.append(r"(?<!\d)\d{12}(?!\d)|(?<!\d)\d{9}(?!\d)")
        elif tok == "email":
            parts.append(EMAIL_RE.pattern)
        elif tok == "phone":
            parts.append(r"(?<!\d)0[\d .]{9,13}(?!\d)")
    if parts:
        return re.compile("|".join(parts))
    if ftype == "date":
        return re.compile(DATE_RE.pattern + "|" + DATE_WORDS_RE.pattern, re.I)
    if ftype == "money":
        return MONEY_RE
    if ftype == "number":
        return re.compile(r"-?\d+(?:[.,]\d+)?")
    return None


# --------------------------------------------------------------------------- #
# POST-PROCESSING: ngày / tiền / số / tên / checkbox / tiếng Việt
# --------------------------------------------------------------------------- #
def _parse_date(text: str) -> Optional[dt.date]:
    t = clean_whitespace(text)
    m = DATE_RE.search(t) or DATE_WORDS_RE.search(t)
    if m:
        d, mo, y = (int(x) for x in m.groups())
    else:
        m2 = re.search(r"(?<!\d)(\d{4})\s*[/\-.]\s*(\d{1,2})\s*[/\-.]\s*(\d{1,2})(?!\d)", t)
        if not m2:
            return None
        y, mo, d = (int(x) for x in m2.groups())
    try:
        if not 1900 <= y <= 2100:
            return None
        return dt.date(y, mo, d)
    except ValueError:
        return None


def normalize_date(text: str) -> Optional[str]:
    """Chuẩn hoá mọi dạng ngày thường gặp về dd/MM/yyyy; None nếu không phải ngày hợp lệ."""
    d = _parse_date(text)
    return d.strftime("%d/%m/%Y") if d else None


_MULTIPLIERS = {"tỷ": 10 ** 9, "tỉ": 10 ** 9, "ty": 10 ** 9, "triệu": 10 ** 6, "trieu": 10 ** 6,
                "nghìn": 10 ** 3, "ngàn": 10 ** 3, "nghin": 10 ** 3, "ngan": 10 ** 3}


def parse_decimal(text: str) -> Optional[Decimal]:
    """Đọc số có thể dùng dấu phẩy thập phân kiểu Việt ('80,5') hoặc dấu chấm."""
    t = re.sub(r"[^\d,.\-]", "", clean_whitespace(text))
    if not t or not re.search(r"\d", t):
        return None
    if "," in t and "." in t:  # dấu xuất hiện sau cùng là thập phân
        dec = "," if t.rfind(",") > t.rfind(".") else "."
        thou = "." if dec == "," else ","
        t = t.replace(thou, "").replace(dec, ".")
    elif "," in t or "." in t:
        sep = "," if "," in t else "."
        if t.count(sep) > 1 or re.fullmatch(rf"-?\d{{1,3}}(?:\{sep}\d{{3}})+", t):
            t = t.replace(sep, "")      # phân tách hàng nghìn
        else:
            t = t.replace(sep, ".")      # thập phân
    try:
        return Decimal(t)
    except InvalidOperation:
        return None


def parse_money_to_int(text: str) -> Optional[int]:
    """'5.200.000.000 đồng' / '5,2 tỷ' / '5 200 000 000' -> 5200000000. Cắt phần '(bằng chữ ...)'."""
    t = clean_whitespace(text)
    t = re.split(r"\(|b[aằ]ng\s*ch[uữ]", t, maxsplit=1, flags=re.I)[0]
    low = t.lower()
    mult = 1
    for word, m in _MULTIPLIERS.items():
        if re.search(rf"\d\s*{word}\b", low):
            mult = m
            break
    t = repair_digits(re.sub(r"(?i)\b(vnd|vnđ|đồng|dong|đ)\b|[₫đ]", "", t))
    num = None
    m = re.fullmatch(r"\s*(\d{1,3}(?:[.,\s]\d{3})+|\d+)\s*", re.sub(r"[^\d.,\s]", "", t))
    if m:
        num = Decimal(re.sub(r"[.,\s]", "", m.group(1)))
    else:
        num = parse_decimal(t)
    if num is None:
        return None
    return int((num * mult).to_integral_value())


def format_money_display(value: int, unit: str = "VND") -> str:
    """5200000000 -> '5.200.000.000 VND' (định dạng hiển thị của đề bài 5.6)."""
    return f"{value:,}".replace(",", ".") + (f" {unit}" if unit else "")


_DIGIT_FIX = str.maketrans({"O": "0", "o": "0", "Q": "0", "D": "0", "I": "1", "l": "1", "|": "1", "!": "1",
                            "S": "5", "s": "5", "B": "8", "Z": "2", "z": "2", "G": "6"})


def repair_digits(text: str) -> str:
    """Sửa nhầm ký tự-số của OCR (O->0, l->1, S->5, B->8...) – chỉ gọi trong ngữ cảnh trường số."""
    def fix_token(tok: str) -> str:
        letters = sum(c.isalpha() for c in tok)
        digits = sum(c.isdigit() for c in tok)
        # chỉ sửa khi token chủ yếu là chữ số (tránh phá từ thường)
        if digits >= 2 and letters <= max(1, digits // 3):
            return tok.translate(_DIGIT_FIX)
        return tok
    return re.sub(r"\S+", lambda m: fix_token(m.group(0)), text)


# --- Số bằng chữ tiếng Việt -> số nguyên (đối chiếu "bằng số" / "bằng chữ") ---
_WORD_DIGITS = {"khong": 0, "mot": 1, "hai": 2, "ba": 3, "bon": 4, "tu": 4, "nam": 5, "lam": 5,
                "sau": 6, "bay": 7, "tam": 8, "chin": 9}
_WORD_SCALES = {"nghin": 10 ** 3, "ngan": 10 ** 3, "trieu": 10 ** 6, "ty": 10 ** 9, "ti": 10 ** 9}
_WORD_SKIP = {"dong", "chan", "vnd", "bang", "chu", "tien", "le", "linh", "va"}


def words_to_int(text: str) -> Optional[int]:
    """
    'Năm tỷ hai trăm triệu đồng' -> 5_200_000_000. Trả về None nếu gặp từ lạ (không đoán).
    Bao phủ: mười/mươi, mốt, tư, lăm, lẻ/linh, trăm, nghìn/ngàn, triệu, tỷ/tỉ.
    Làm việc trên dạng đã bỏ dấu nên chịu được OCR làm rơi dấu.
    """
    t = re.sub(r"\bviet nam\b", "", fold_vietnamese(text))
    total, group, pending, saw_number = 0, 0, None, False
    for tok in t.split():
        if tok in _WORD_SKIP:
            continue
        if tok in _WORD_DIGITS:
            pending, saw_number = _WORD_DIGITS[tok], True
        elif tok == "muoi":                       # "mười" = 10, "hai mươi" = 20
            group += (pending if pending is not None else 1) * 10
            pending, saw_number = None, True
        elif tok == "tram":
            group += (pending if pending is not None else 1) * 100
            pending, saw_number = None, True
        elif tok in _WORD_SCALES:
            if pending is not None:
                group, pending = group + pending, None
            total += (group if group else 1) * _WORD_SCALES[tok]
            group, saw_number = 0, True
        else:
            return None
    if pending is not None:
        group += pending
    return total + group if saw_number else None


# --- Tiếng Việt: sửa dấu / proper-case ---
# Các cụm chuẩn hay gặp trong hồ sơ ngân hàng – khi OCR làm mất/sai dấu sẽ được khôi phục.
_GLOSSARY = [
    "Thành phố Hồ Chí Minh", "Cộng hòa xã hội chủ nghĩa Việt Nam", "Độc lập - Tự do - Hạnh phúc",
    "Ngân hàng Thương mại Cổ phần", "Văn phòng Công chứng", "Công chứng viên", "Quyền sử dụng đất",
    "Giấy chứng nhận", "Sở Tài nguyên và Môi trường", "Văn phòng Đăng ký đất đai", "Hợp đồng bảo đảm",
    "Hợp đồng thế chấp", "Bên thế chấp", "Bên nhận thế chấp", "Bên bảo đảm", "Bên nhận bảo đảm",
    "Tài sản gắn liền với đất", "Đất ở tại đô thị", "Lâu dài", "Nhà ở riêng lẻ", "Phường", "Quận",
    "Thành phố", "Tỉnh", "Huyện", "Xã", "Thị trấn", "Việt Nam", "Giám đốc", "Phòng Giao dịch", "Chi nhánh",
    "Công ty TNHH", "Thẩm định giá", "Bổ sung vốn lưu động", "Bản chính", "Bản sao", "Nguyễn", "Trần", "Lê",
    "Phạm", "Hoàng", "Huỳnh", "Phan", "Vũ", "Võ", "Đặng", "Bùi", "Đỗ", "Hồ", "Ngô", "Dương", "Lý",
    "Bến Thành", "Bến Nghé", "Đa Kao", "Điện Biên Phủ", "Lê Lợi", "Nguyễn Trãi",
    "Số", "Thửa đất số", "Tờ bản đồ số", "Phòng Công chứng", "Kho tài sản bảo đảm",
]
_GLOSSARY_BY_LEN: Dict[int, Dict[str, str]] = {}
for _g in _GLOSSARY:
    _GLOSSARY_BY_LEN.setdefault(len(_g.split()), {})[fold_vietnamese(_g)] = _g

# lỗi dấu phổ biến: tổ hợp dấu bị tách rời / ký tự lạ thay dấu
_CHAR_FIXES = {"\u0300\u0301": "\u0301", "’": "'", "‘": "'", "“": '"', "”": '"', "–": "-", "—": "-", "…": "..."}
_TONE_STYLE_FIXES = [  # kiểu bỏ dấu cũ -> mới (oà->òa, uý->úy...) để thống nhất khi so sánh
    ("hoà", "hòa"), ("Hoà", "Hòa"), ("HOÀ", "HÒA"), ("hoá", "hóa"), ("Hoá", "Hóa"), ("HOÁ", "HÓA"),
    ("thuỷ", "thủy"), ("Thuỷ", "Thủy"), ("quỳ", "quỳ"), ("uý", "úy"), ("Uý", "Úy"),
]


def _match_case(src: str, dst: str) -> str:
    if src.isupper():
        return dst.upper()
    if src.islower():
        return dst.lower()
    if src.istitle():
        return " ".join(w[:1].upper() + w[1:] for w in dst.split())
    return dst


def fix_vietnamese_text(text: str) -> str:
    """
    Sửa lỗi tiếng Việt thường gặp sau OCR:
      * chuẩn hoá Unicode NFC, gộp dấu tách rời, ký tự nháy/gạch lạ;
      * thống nhất kiểu đặt dấu (hoà -> hòa);
      * khôi phục dấu cho các cụm chuẩn (địa danh, chức danh, thuật ngữ ngân hàng) khi OCR làm
        mất dấu hoặc sai dấu nhẹ (so khớp mờ >= 92 trên dạng bỏ dấu).
    """
    t = nfc(text)
    for a, b in _CHAR_FIXES.items():
        t = t.replace(a, b)
    for a, b in _TONE_STYLE_FIXES:
        t = t.replace(a, b)
    words = t.split(" ")
    out: List[str] = []
    i = 0
    max_n = max(_GLOSSARY_BY_LEN)
    while i < len(words):
        replaced = False
        for n in range(min(max_n, len(words) - i), 0, -1):
            window = words[i:i + n]
            core = [re.sub(r"^[^\w]+|[^\w]+$", "", w) for w in window]
            if not all(core):
                continue
            key = fold_vietnamese(" ".join(core))
            table = _GLOSSARY_BY_LEN.get(n, {})
            cand = table.get(key)
            if cand is None and n >= 2:
                best = max(table.items(), key=lambda kv: fuzz.ratio(kv[0], key), default=None)
                if best and fuzz.ratio(best[0], key) >= 92:
                    cand = best[1]
            if cand is not None:
                joined = " ".join(core)
                fixed = _match_case(joined, cand) if (joined.isupper() or joined.islower()) else cand
                prefix = re.match(r"^[^\w]*", window[0]).group(0)
                suffix = re.search(r"[^\w]*$", window[-1]).group(0)
                out.append(prefix + fixed + suffix)
                i += n
                replaced = True
                break
        if not replaced:
            out.append(words[i])
            i += 1
    return " ".join(out)


def proper_case_vn(text: str) -> str:
    """'TRẦN VĂN MINH' / 'trần văn minh' -> 'Trần Văn Minh'."""
    return " ".join(w[:1].upper() + w[1:].lower() for w in clean_whitespace(text).split(" ") if w)


# --- Checkbox ---
_CHECK_MARKS = r"(?:☑|☒|✓|✔|✗|✘|\[\s*[xXvV✓]\s*\]|\(\s*[xXvV✓]\s*\)|(?<![a-z])[xXvV](?![a-z]))"
_TRUE_WORDS = {"true", "yes", "co", "x", "checked", "dung", "1"}
_FALSE_WORDS = {"false", "no", "khong", "unchecked", "sai", "0"}


def _fold_keep_marks(text: str) -> str:
    t = strip_diacritics(nfc(text)).lower()
    return re.sub(r"\s+", " ", t)


def normalize_checkbox(spec: Any, raw: str) -> str:
    """
    * Có `enum:a|b` -> trả về đúng 1 giá trị enum (so khớp nhãn 'to_chuc' ~ 'tổ chức'; nếu cả hai
      nhãn đều xuất hiện thì chọn nhãn đứng ngay sau dấu tích ☑/[x]).
    * Không có enum -> 'true'/'false'.
    """
    text = clean_whitespace(raw)
    if not text:
        return ""
    enums = spec.enum_values
    folded = _fold_keep_marks(text)
    if enums:
        for e in enums:                       # đã đúng giá trị enum
            if folded.strip() == _fold_keep_marks(e):
                return e
        labels = {e: _fold_keep_marks(e.replace("_", " ")) for e in enums}
        present = [e for e, lab in labels.items() if lab and lab in folded]
        if len(present) == 1:
            return present[0]
        marked = []
        for e in present:
            if re.search(_CHECK_MARKS + r"\s*" + re.escape(labels[e]), folded, flags=re.I):
                marked.append(e)
        if len(marked) == 1:
            return marked[0]
        best = max(enums, key=lambda e: fuzz.partial_ratio(labels[e], folded))
        return best if fuzz.partial_ratio(labels[best], folded) >= 85 and len(present) <= 1 else text
    if re.search(_CHECK_MARKS, text) or folded.strip() in _TRUE_WORDS:
        return "true"
    if re.search(r"☐|□|\[\s*\]|\(\s*\)", text) or folded.strip() in _FALSE_WORDS:
        return "false"
    return text


# --- Điều phối post-process theo type ---
@dataclass
class PostResult:
    value: str
    changes: List[str] = field(default_factory=list)
    ok: bool = True           # False nếu không chuẩn hoá được theo type (vd: ngày vô lý)


def _repair_for_validation(validation: str, value: str, ftype: str) -> Tuple[str, Optional[str]]:
    """Thử vài cách sửa 'rẻ' (bỏ dấu phân cách, sửa nhầm O/0...) nếu giá trị gốc không qua validator."""
    if not validation or validate_value(validation, value, ftype).passed:
        return value, None
    low = validation.lower()
    if low.startswith("enum:"):
        return value, None
    candidates = [("bỏ ký tự phân cách", re.sub(r"[\s.\-_/]", "", value)),
                  ("sửa nhầm ký tự-số", repair_digits(value)),
                  ("sửa nhầm ký tự-số + bỏ phân cách", re.sub(r"[\s.\-_/]", "", repair_digits(value)))]
    for why, cand in candidates:
        if cand != value and validate_value(validation, cand, ftype).passed:
            return cand, why
    return value, None


def postprocess_value(spec: Any, raw: Optional[str], keep_verbatim: bool = False) -> PostResult:
    """
    Chuẩn hoá một giá trị thô theo `spec.type` (FieldSpec hoặc ColumnSpec).

    keep_verbatim=True: chỉ làm sạch khoảng trắng/Unicode (file "Lưu ý" yêu cầu bóc nguyên văn);
    chuẩn hoá đầy đủ khi False (đề bài mục 5.6).
    """
    if raw is None:
        return PostResult("", [], ok=True)
    text = clean_whitespace(str(raw))
    changes: List[str] = []
    if not text:
        return PostResult("", [])
    ftype = spec.type

    if keep_verbatim:
        return PostResult(text, [])

    if ftype == "date":
        norm = normalize_date(text)
        if norm is None:
            return PostResult(text, ["không phải ngày hợp lệ"], ok=False)
        if norm != text:
            changes.append("ngày -> dd/MM/yyyy")
        return PostResult(norm, changes)

    if ftype == "money":
        val = parse_money_to_int(text)
        if val is None:
            return PostResult(text, ["không đọc được số tiền"], ok=False)
        norm = str(val)
        if norm != text:
            changes.append("số tiền -> số nguyên")
        return PostResult(norm, changes)

    if ftype == "number":
        fixed = repair_digits(text)
        if fixed != text:
            changes.append("sửa nhầm ký tự-số")
        m = re.search(r"-?\d[\d.,\s]*", fixed)
        if not m:
            return PostResult(text, ["không phải số"], ok=False)
        norm = re.sub(r"\s+", "", m.group(0)).rstrip(".,")
        if spec.validation.startswith("^") and re.fullmatch(r"\^?\\d[+*]?\$?", spec.validation):
            norm = re.sub(r"\D", "", norm)       # validation chỉ chấp nhận chữ số
        return PostResult(norm, changes + (["tách số khỏi đơn vị"] if norm != fixed.strip() else []))

    if ftype in ("checkbox", "boolean"):
        norm = normalize_checkbox(spec, text)
        if norm != text:
            changes.append("checkbox -> enum/boolean")
        return PostResult(norm, changes)

    # text
    fixed = fix_vietnamese_text(text)
    if fixed != text:
        changes.append("sửa dấu/Unicode tiếng Việt")
    text = fixed
    if spec.wants_proper_case:
        pc = proper_case_vn(text)
        if pc != text:
            changes.append("proper-case")
        text = pc
    text, why = _repair_for_validation(spec.validation, text, ftype)
    if why:
        changes.append(why)
    return PostResult(text, changes)


# --------------------------------------------------------------------------- #
# Khoá so sánh dùng chung (Fusion + Eval)
# --------------------------------------------------------------------------- #
def comparison_key(ftype: str, value: Optional[str], enum_values: Optional[Iterable[str]] = None) -> str:
    """
    Khoá so sánh sau chuẩn hoá: bỏ khoảng trắng thừa, lowercase, đồng bộ ngày/số tiền/số.
    Hai giá trị 'giống nhau về nghĩa' có cùng khoá.
    """
    v = clean_whitespace(value or "")
    if not v:
        return ""
    if ftype == "date":
        return normalize_date(v) or v.lower()
    if ftype == "money":
        m = parse_money_to_int(v)
        return str(m) if m is not None else re.sub(r"\D", "", v)
    if ftype == "number":
        d = parse_decimal(repair_digits(v))
        if d is None:
            return v.lower()
        s = format(d.normalize(), "f")
        return s
    if ftype in ("checkbox", "boolean"):
        return v.lower()
    if re.fullmatch(r"[\d\s.\-/]+", v) and sum(c.isdigit() for c in v) >= 6:
        return re.sub(r"\D", "", v)          # SĐT/CCCD/MST: '0909 123 456' == '0909123456'
    t = fix_vietnamese_text(v).casefold()
    t = re.sub(r"\s+", " ", t).strip(" .,;:-_")
    return t
